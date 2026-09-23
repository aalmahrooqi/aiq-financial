# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0

"""Read XLSX grids into self-contained table chunks, independent of NeMo APIs."""

from __future__ import annotations

import html
import re
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date
from datetime import datetime
from datetime import time
from datetime import timedelta
from math import ceil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

MAX_CHUNK_TOKENS = 3000
_MAX_SHEET_CELLS = 2_000_000
_SHEET_CONTEXT_ROWS = 8
_TOKEN_PATTERN = re.compile(r"\w+|[^\w\s]", re.UNICODE)


def estimate_token_count(text: str) -> int:
    """Conservatively estimate BPE tokens without downloading a model tokenizer."""
    count = text.count("\n")
    for token in _TOKEN_PATTERN.findall(text):
        count += max(1, ceil(len(token.encode("utf-8")) / 3)) if token[0].isalnum() or token[0] == "_" else 1
    return count


@dataclass(frozen=True)
class ExcelChunk:
    """Searchable text plus original cell data and workbook coordinates."""

    text: str
    metadata: dict[str, Any]


def _cell_data(cell: Any, cached: Any) -> dict[str, Any]:
    formula = cell.value if cell.data_type == "f" else None
    value = cached.value if formula is not None else cell.value
    if isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    elif isinstance(value, timedelta):
        value = str(value)
    result = {"column": cell.column_letter, "value": value}
    if formula is not None:
        result["formula"] = formula if isinstance(formula, str) else getattr(formula, "text", "[data table formula]")
    if cell.number_format != "General":
        result["number_format"] = cell.number_format
    return result


def _cell_text(cell: dict[str, Any]) -> str:
    value = cell["value"]
    if value is None:
        text = f"[no cached value: {cell['formula']}]" if "formula" in cell else ""
    elif isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, (int, float)):
        number_format = re.sub(r'"[^"]*"|\\.', "", cell.get("number_format", ""))
        text = f"{value * 100:.15g}%" if "%" in number_format else f"{value:.15g}"
    else:
        text = str(value)
    return html.escape(text).replace("|", "&#124;").replace("\r\n", "\n").replace("\n", "<br>")


def _read_rows(sheet: Worksheet, cached_sheet: Any) -> dict[int, dict[int, dict[str, Any]]]:
    """Keep populated cells only; formatting-only rows do not become content."""
    if sheet.max_row * sheet.max_column > _MAX_SHEET_CELLS:
        raise ValueError(f"Excel sheet {sheet.title!r} exceeds the supported used-range size")
    rows = {}
    cached_rows = cached_sheet.iter_rows(max_row=sheet.max_row, max_col=sheet.max_column)
    for cells, cached_cells in zip(sheet.iter_rows(), cached_rows, strict=True):
        populated = {
            cell.column: _cell_data(cell, cached)
            for cell, cached in zip(cells, cached_cells, strict=True)
            if cell.value is not None and (not isinstance(cell.value, str) or cell.value.strip())
        }
        if populated:
            rows[cells[0].row] = populated
    return rows


def _row_regions(rows: dict[int, dict[int, dict[str, Any]]]) -> Iterator[str]:
    """Group consecutive populated rows, preserving side-by-side column layouts."""
    populated = [row for row, cells in rows.items() if cells]
    start = 0
    while start < len(populated):
        end = start + 1
        while end < len(populated) and populated[end] == populated[end - 1] + 1:
            end += 1
        block = populated[start:end]
        columns = [col for row in block for col in rows[row]]
        left, right = min(columns), max(columns)
        yield f"{get_column_letter(left)}{block[0]}:{get_column_letter(right)}{block[-1]}"
        start = end


def _table_chunks(
    sheet: Worksheet,
    rows: dict[int, dict[int, dict[str, Any]]],
    region: str,
    table_name: str | None = None,
    header_count: int = 0,
) -> Iterator[ExcelChunk]:
    left, top, right, bottom = range_boundaries(region)
    if (right - left + 1) * (bottom - top + 1) > _MAX_SHEET_CELLS:
        raise ValueError(f"Excel table {region!r} exceeds the supported used-range size")
    # Opening rows commonly contain units and year headers shared by several
    # sections. Their original row numbers distinguish context from new data.
    context = [row for row in rows if rows[row] and row <= _SHEET_CONTEXT_ROWS and row < top]
    context += [
        merged.min_row
        for merged in sheet.merged_cells.ranges
        if merged.min_row < merged.max_row
        and merged.min_row <= bottom
        and merged.max_row >= top
        and merged.min_col <= right
        and merged.max_col >= left
        and rows.get(merged.min_row)
    ]
    headers = list(range(top, min(top + header_count, bottom + 1)))
    if not table_name:
        for row in range(top, min(top + 3, bottom + 1)):
            cells = [cell for col, cell in rows.get(row, {}).items() if left <= col <= right]
            if not cells or any(isinstance(cell["value"], (int, float)) or "formula" in cell for cell in cells):
                break
            headers.append(row)
    repeated = sorted(set(context + headers))
    columns = sorted(set(range(left, right + 1)).union(*(set(rows[row]) for row in context)))
    merges = [
        str(merged)
        for merged in sorted(sheet.merged_cells.ranges, key=str)
        if merged.min_col <= max(columns)
        and merged.max_col >= min(columns)
        and (
            (merged.min_row <= bottom and merged.max_row >= top)
            or any(merged.min_row <= row <= merged.max_row for row in repeated)
        )
    ]

    def selected_cells(row: int) -> dict[int, dict[str, Any]]:
        return {col: cell for col, cell in rows.get(row, {}).items() if row in context or left <= col <= right}

    def line(row: int) -> str:
        cells = selected_cells(row)
        values = [_cell_text(cells[col]) if col in cells else "" for col in columns]
        return "| " + " | ".join([str(row), *values]) + " |"

    prefix = f"Sheet: {sheet.title}\nRegion: {region}\n"
    if table_name:
        prefix += f"Table: {table_name}\n"
    if merges:
        prefix += f"Merged cells (value at top-left): {', '.join(merges)}\n"
    if repeated:
        prefix += f"Context/header rows repeated in each chunk: {', '.join(map(str, repeated))}\n"
    prefix += "| Row | " + " | ".join(get_column_letter(col) for col in columns) + " |\n"
    prefix += "| --- |" + " --- |" * len(columns) + "\n"
    prefix += "".join(line(row) + "\n" for row in repeated)

    def chunk(numbers: list[int]) -> ExcelChunk:
        selected = sorted(set(repeated + numbers))
        cell_range = f"{get_column_letter(left)}{numbers[0]}:{get_column_letter(right)}{numbers[-1]}"
        return ExcelChunk(
            text=prefix + "\n".join(line(row) for row in numbers if row not in repeated),
            metadata={
                "sheet_name": sheet.title,
                "sheet_state": sheet.sheet_state,
                "cell_range": cell_range,
                "table_range": region,
                "table_name": table_name,
                "structured_data": {
                    "columns": [get_column_letter(col) for col in columns],
                    "context_rows": repeated,
                    "merged_cells": merges,
                    "rows": [{"row": row, "cells": list(selected_cells(row).values())} for row in selected],
                },
            },
        )

    batch: list[int] = []
    prefix_tokens = estimate_token_count(prefix)
    size = prefix_tokens
    for row in range(top, bottom + 1):
        row_size = 0 if row in repeated else estimate_token_count(line(row)) + 1
        if prefix_tokens + row_size > MAX_CHUNK_TOKENS:
            raise ValueError(
                f"Excel sheet {sheet.title!r}, row {row} exceeds the {MAX_CHUNK_TOKENS}-token indexing budget"
            )
        if batch and size + row_size > MAX_CHUNK_TOKENS:
            yield chunk(batch)
            batch = []
            size = prefix_tokens
        batch.append(row)
        size += row_size
    if batch:
        yield chunk(batch)


def read_excel_chunks(path: Path) -> list[ExcelChunk]:
    """Read all worksheets using saved formula results, without modifying the file.

    Declared Excel tables are preferred. Other populated cells are preserved as
    row regions, including blank columns, rather than inferring relationships
    between neighboring tables. Empty worksheets are skipped.
    """
    workbook = load_workbook(path, data_only=False, keep_links=False)
    try:
        cached = load_workbook(path, data_only=True, read_only=True, keep_links=False)
        try:
            chunks = []
            for sheet in workbook.worksheets:
                rows = _read_rows(sheet, cached[sheet.title])
                remaining = {row: dict(cells) for row, cells in rows.items()}
                for table in sheet.tables.values():
                    chunks.extend(_table_chunks(sheet, rows, table.ref, table.name, table.headerRowCount or 0))
                    left, top, right, bottom = range_boundaries(table.ref)
                    for row in range(top, bottom + 1):
                        for col in range(left, right + 1):
                            remaining.get(row, {}).pop(col, None)
                for region in _row_regions(remaining):
                    chunks.extend(_table_chunks(sheet, remaining, region))
            if not chunks:
                raise ValueError("Excel workbook contains no populated cells")
            return chunks
        finally:
            cached.close()
    finally:
        workbook.close()
